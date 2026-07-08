import io
from datetime import date, timedelta

import openpyxl
from django.test import TestCase
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APIClient

from curriculum.models import Filiere, Module, Semestre
from users.models import Amphitheater, Block, Classe, Room, Student, User
from courses.models import Session
from .models import Attendance


class ExportWizardTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        self.admin = User.objects.create_user(
            email="admin@test.ma",
            password="password123",
            first_name="Admin",
            last_name="Root",
            role="admin",
        )
        self.teacher = User.objects.create_user(
            email="teacher@test.ma",
            password="password123",
            first_name="Jean",
            last_name="Dupont",
            role="teacher",
        )
        self.teacher2 = User.objects.create_user(
            email="teacher2@test.ma",
            password="password123",
            first_name="Marie",
            last_name="Curie",
            role="teacher",
        )

        self.semestre = Semestre.objects.create(name="S5")
        self.filiere = Filiere.objects.create(
            name="Genie Informatique", semestre=self.semestre, niveau="LIC"
        )
        self.module = Module.objects.create(
            name="Programmation Web", filiere=self.filiere
        )

        self.classe = Classe.objects.create(
            name="GI S5",
            filiere=self.filiere,
            semestre=self.semestre,
            module=self.module,
            niveau="LIC",
            academic_year="2025-2026",
        )

        self.student1 = Student.objects.create(
            first_name="Ahmed",
            last_name="Benali",
            code_massar="CNE001",
            code_apogee="APO001",
            email="ahmed@test.ma",
            classe=self.classe,
        )
        self.student2 = Student.objects.create(
            first_name="Sara",
            last_name="El Idrissi",
            code_massar="CNE002",
            code_apogee="APO002",
            email="sara@test.ma",
            classe=self.classe,
        )

        self.session1 = Session.objects.create(
            subject="Programmation Web",
            module=self.module,
            teacher=self.teacher,
            classe=self.classe,
            date=date.today() - timedelta(days=5),
            start_time="08:00",
            end_time="10:00",
        )
        self.session2 = Session.objects.create(
            subject="Programmation Web",
            module=self.module,
            teacher=self.teacher,
            classe=self.classe,
            date=date.today() - timedelta(days=3),
            start_time="10:00",
            end_time="12:00",
        )

        self.att1 = Attendance.objects.create(
            student=self.student1,
            session=self.session1,
            attendance_status=Attendance.STATUS_CONFIRMED,
        )
        self.att2 = Attendance.objects.create(
            student=self.student2,
            session=self.session1,
            attendance_status=Attendance.STATUS_ABSENT,
        )
        self.att3 = Attendance.objects.create(
            student=self.student1,
            session=self.session2,
            attendance_status=Attendance.STATUS_CONFIRMED,
        )
        self.att4 = Attendance.objects.create(
            student=self.student2,
            session=self.session2,
            attendance_status=Attendance.STATUS_CONFIRMED,
        )

    def _login(self, user):
        self.client.force_authenticate(user=user)

    def _export_url(self):
        return "/api/v1/attendance/export-wizard/"

    # ── Admin tests ──

    def test_admin_export_all(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_type": "all",
                "export_mode": "by_student",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertIsNotNone(ws)
        self.assertGreater(ws.max_row, 1)

    def test_admin_export_present_only(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_type": "present",
                "export_mode": "by_student",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[-1], "Présent")

    def test_admin_export_absent_only(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_type": "absent",
                "export_mode": "by_student",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[-1], "Absent")

    def test_admin_export_by_student(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_mode": "by_student",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Code Apogée", headers)
        self.assertIn("Code Massar / CNE", headers)

    def test_admin_export_by_module(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_mode": "by_module",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_admin_export_by_session(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_mode": "by_session",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_admin_export_resume_absences(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_mode": "resume_absences",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Nombre total de séances", headers)
        self.assertIn("Taux de présence (%)", headers)

    def test_admin_export_resume_presences(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_mode": "resume_presences",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_admin_export_with_teacher_filter(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_mode": "by_student",
                "enseignant_id": self.teacher.id,
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_admin_export_with_filters(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "semestre": "S5",
                "filiere": "Genie Informatique",
                "module": "Programmation Web",
                "classe_id": self.classe.id,
                "export_type": "all",
                "export_mode": "by_student",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    # ── Teacher tests ──

    def test_teacher_export_own_sessions(self):
        self._login(self.teacher)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_teacher_cannot_see_other_teacher_data(self):
        Session.objects.create(
            subject="Other Module",
            module=self.module,
            teacher=self.teacher2,
            classe=self.classe,
            date=date.today() - timedelta(days=1),
            start_time="08:00",
            end_time="10:00",
        )
        self._login(self.teacher)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        row_count = ws.max_row - 1
        self.assertEqual(row_count, 4)

    def test_teacher_export_present_only(self):
        self._login(self.teacher)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_type": "present",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[-1], "Présent")

    def test_teacher_export_absent_only(self):
        self._login(self.teacher)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_type": "absent",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[-1], "Absent")

    def test_teacher_export_resume(self):
        self._login(self.teacher)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
                "export_mode": "resume_presences",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Nombre total de séances", headers)

    # ── Validation tests ──

    def test_requires_dates(self):
        self._login(self.admin)
        resp = self.client.get(self._export_url(), {})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_invalid_export_mode(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": "2025-01-01",
                "end_date": "2025-12-31",
                "export_mode": "invalid_mode",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_invalid_export_type(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": "2025-01-01",
                "end_date": "2025-12-31",
                "export_type": "invalid_type",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_excel_format_valid(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
            },
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resp["Content-Type"],
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIsNotNone(wb)

    def test_excel_headers_correct(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
            },
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        expected = [
            "Code Apogée",
            "Code Massar / CNE",
            "Nom",
            "Prénom",
            "Email",
            "Filière",
            "Classe",
            "Semestre",
            "Module",
            "Enseignant",
            "Date de la séance",
            "Heure de début",
            "Heure de fin",
            "Heure du scan",
            "Statut",
        ]
        self.assertEqual(headers, expected)

    def test_excel_contains_exported_data(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
            },
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        codes = [r[1] for r in data_rows]
        self.assertIn("CNE001", codes)
        self.assertIn("CNE002", codes)
        self.assertIn("APO001", [r[0] for r in data_rows])

    # ── Permission tests ──

    def test_student_no_access(self):
        student_user = User.objects.create_user(
            email="student@test.ma",
            password="password123",
            first_name="Student",
            last_name="Test",
            role="student",
        )
        self._login(student_user)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": "2025-01-01",
                "end_date": "2025-12-31",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_unauthenticated_no_access(self):
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": "2025-01-01",
                "end_date": "2025-12-31",
            },
        )
        self.assertEqual(resp.status_code, status.HTTP_401_UNAUTHORIZED)

    # ── File download tests ──

    def test_download_filename_has_extension(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
            },
        )
        self.assertIn("Content-Disposition", resp)
        self.assertIn(".xlsx", resp["Content-Disposition"])

    # ── Student filter tests ──

    def test_export_single_student_no_dates(self):
        self._login(self.admin)
        resp = self.client.get(self._export_url(), {"etudiant_id": self.student1.id})
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resp["Content-Type"],
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        codes = [r[1] for r in data_rows if r[1]]
        self.assertIn("CNE001", codes)
        self.assertNotIn("CNE002", codes)

    def test_export_single_student_summary_appended(self):
        self._login(self.admin)
        resp = self.client.get(self._export_url(), {"etudiant_id": self.student1.id})
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        cell_values = [cell for row in all_rows for cell in (row or []) if cell]
        self.assertIn("RÉSUMÉ", cell_values)
        self.assertIn("Nombre total de séances", cell_values)
        self.assertIn("Nombre de présences", cell_values)
        self.assertIn("Taux de présence (%)", cell_values)
        self.assertIn(2, cell_values)
        self.assertIn(2, cell_values)

    def test_export_single_student_no_dates_teacher(self):
        self._login(self.teacher)
        resp = self.client.get(self._export_url(), {"etudiant_id": self.student1.id})
        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        codes = [r[1] for r in data_rows if r[1]]
        self.assertIn("CNE001", codes)

    def test_export_dates_still_required_without_student(self):
        self._login(self.admin)
        resp = self.client.get(self._export_url(), {})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("dates", resp.data.get("error", "").lower())

    def test_export_invalid_student_id(self):
        self._login(self.admin)
        resp = self.client.get(self._export_url(), {"etudiant_id": "abc"})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_export_single_student_with_dates(self):
        self._login(self.admin)
        resp = self.client.get(
            self._export_url(),
            {
                "etudiant_id": self.student1.id,
                "start_date": (date.today() - timedelta(days=30)).isoformat(),
                "end_date": date.today().isoformat(),
            },
        )
        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        codes = [r[1] for r in data_rows if r[1]]
        self.assertIn("CNE001", codes)
        self.assertNotIn("CNE002", codes)


class RoomConflictTests(TestCase):
    """Tests that two sessions cannot share the same room/amphi at the same time."""

    def setUp(self):
        self.client = APIClient()

        self.admin = User.objects.create_user(
            email="admin@test.ma", password="password123",
            first_name="Admin", last_name="Root", role="admin",
        )
        self.teacher = User.objects.create_user(
            email="teacher@test.ma", password="password123",
            first_name="Jean", last_name="Dupont", role="teacher",
        )
        self.teacher2 = User.objects.create_user(
            email="teacher2@test.ma", password="password123",
            first_name="Marie", last_name="Curie", role="teacher",
        )

        self.semestre = Semestre.objects.create(name="S5")
        self.filiere = Filiere.objects.create(
            name="Genie Informatique", semestre=self.semestre, niveau="LIC"
        )
        self.module = Module.objects.create(
            name="Programmation Web", filiere=self.filiere
        )
        self.classe = Classe.objects.create(
            name="GI S5", filiere=self.filiere, semestre=self.semestre,
            module=self.module, niveau="LIC", academic_year="2025-2026",
        )

        self.block = Block.objects.create(name="Bloc A")
        self.room1 = Room.objects.create(name="Salle 101", block=self.block)
        self.room2 = Room.objects.create(name="Salle 102", block=self.block)
        self.amphi = Amphitheater.objects.create(name="Amphi 1")

        self.today = date.today()

    def _login(self, user):
        self.client.force_authenticate(user=user)

    def _session_url(self):
        return "/api/v1/sessions/"

    def _available_url(self):
        return "/api/v1/sessions/available-rooms/"

    def test_two_teachers_cannot_book_same_room_same_time(self):
        self._login(self.admin)
        resp = self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        resp2 = self.client.post(self._session_url(), {
            "subject": "Physique", "module_id": self.module.id,
            "teacher": self.teacher2.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        self.assertEqual(resp2.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("occup", str(resp2.data).lower())

    def test_room_available_when_free(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        resp = self.client.post(self._session_url(), {
            "subject": "Physique", "module_id": self.module.id,
            "teacher": self.teacher2.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room2.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)

    def test_edit_session_ignores_self(self):
        self._login(self.admin)
        resp = self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        sid = resp.data["id"]
        resp2 = self.client.patch(
            self._session_url() + str(sid) + "/",
            {"subject": "Maths Avancees"},
            format="json",
        )
        self.assertEqual(resp2.status_code, status.HTTP_200_OK)

    def test_cannot_edit_into_conflicting_room(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        resp2 = self.client.post(self._session_url(), {
            "subject": "Physique", "module_id": self.module.id,
            "teacher": self.teacher2.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room2.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        self.assertEqual(resp2.status_code, status.HTTP_201_CREATED)
        sid2 = resp2.data["id"]
        resp3 = self.client.patch(
            self._session_url() + str(sid2) + "/",
            {"room": self.room1.id},
            format="json",
        )
        self.assertEqual(resp3.status_code, status.HTTP_400_BAD_REQUEST)

    def test_amphitheater_conflict(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Conf", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "amphitheater",
            "amphitheater": self.amphi.id,
            "date": self.today.isoformat(),
            "start_time": "14:00", "end_time": "16:00",
            "session_type": "Cours",
        })
        resp = self.client.post(self._session_url(), {
            "subject": "Sem", "module_id": self.module.id,
            "teacher": self.teacher2.id, "classe": self.classe.id,
            "location_type": "amphitheater",
            "amphitheater": self.amphi.id,
            "date": self.today.isoformat(),
            "start_time": "14:00", "end_time": "16:00",
            "session_type": "Cours",
        })
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_amphitheater_available_different_time(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Conf", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "amphitheater",
            "amphitheater": self.amphi.id,
            "date": self.today.isoformat(),
            "start_time": "14:00", "end_time": "16:00",
            "session_type": "Cours",
        })
        resp = self.client.post(self._session_url(), {
            "subject": "Sem", "module_id": self.module.id,
            "teacher": self.teacher2.id, "classe": self.classe.id,
            "location_type": "amphitheater",
            "amphitheater": self.amphi.id,
            "date": self.today.isoformat(),
            "start_time": "16:00", "end_time": "18:00",
            "session_type": "Cours",
        })
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)

    def test_available_rooms_endpoint_excludes_occupied(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._available_url(), {
            "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "location_type": "block",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        room_ids = [r["id"] for b in resp.data.get("blocks", []) for r in b.get("rooms", [])]
        self.assertNotIn(self.room1.id, room_ids)
        self.assertIn(self.room2.id, room_ids)

    def test_available_endpoint_without_location_type(self):
        self._login(self.admin)
        resp = self.client.get(self._available_url(), {
            "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("blocks", resp.data)
        self.assertIn("amphitheaters", resp.data)

    # ── New available-locations endpoint tests ──

    def _locations_url(self):
        return "/api/v1/sessions/available-locations/"

    def test_available_locations_returns_all_keys(self):
        self._login(self.admin)
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("rooms", resp.data)
        self.assertIn("blocks", resp.data)
        self.assertIn("amphitheaters", resp.data)

    def test_available_locations_requires_params(self):
        self._login(self.admin)
        resp = self.client.get(self._locations_url(), {})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_available_locations_excludes_occupied_room(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        room_ids = [r["id"] for r in resp.data["rooms"]]
        self.assertNotIn(self.room1.id, room_ids)
        self.assertIn(self.room2.id, room_ids)

    def test_available_locations_excludes_occupied_amphitheater(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Conf", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "amphitheater",
            "amphitheater": self.amphi.id,
            "date": self.today.isoformat(),
            "start_time": "14:00", "end_time": "16:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "14:00", "end_time": "16:00",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        amphi_ids = [a["id"] for a in resp.data["amphitheaters"]]
        self.assertNotIn(self.amphi.id, amphi_ids)

    def test_available_locations_blocks_show_available_rooms_only(self):
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        block = next((b for b in resp.data["blocks"] if b["id"] == self.block.id), None)
        self.assertIsNotNone(block)
        room_ids_in_block = [r["id"] for r in block["rooms"]]
        self.assertNotIn(self.room1.id, room_ids_in_block)
        self.assertIn(self.room2.id, room_ids_in_block)

    def test_available_locations_partial_overlap_start(self):
        """Existing 09:00-11:00, new 08:30-09:30 should detect conflict."""
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Existing", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "09:00", "end_time": "11:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "08:30", "end_time": "09:30",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        room_ids = [r["id"] for r in resp.data["rooms"]]
        self.assertNotIn(self.room1.id, room_ids)

    def test_available_locations_partial_overlap_end(self):
        """Existing 09:00-11:00, new 10:30-11:30 should detect conflict."""
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Existing", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "09:00", "end_time": "11:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "10:30", "end_time": "11:30",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        room_ids = [r["id"] for r in resp.data["rooms"]]
        self.assertNotIn(self.room1.id, room_ids)

    def test_available_locations_middle_overlap(self):
        """Existing 09:00-11:00, new 09:15-10:15 should detect conflict."""
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Existing", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "09:00", "end_time": "11:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "09:15", "end_time": "10:15",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        room_ids = [r["id"] for r in resp.data["rooms"]]
        self.assertNotIn(self.room1.id, room_ids)

    def test_available_locations_total_overlap(self):
        """Existing 09:00-11:00, new 08:00-12:00 should detect conflict."""
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Existing", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "09:00", "end_time": "11:00",
            "session_type": "Cours",
        })
        resp = self.client.get(self._locations_url(), {
            "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "12:00",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        room_ids = [r["id"] for r in resp.data["rooms"]]
        self.assertNotIn(self.room1.id, room_ids)

    def test_available_locations_no_conflict_different_date(self):
        """Same room, same time but different date should show as available."""
        self._login(self.admin)
        self.client.post(self._session_url(), {
            "subject": "Maths", "module_id": self.module.id,
            "teacher": self.teacher.id, "classe": self.classe.id,
            "location_type": "block", "block": self.block.id,
            "room": self.room1.id, "date": self.today.isoformat(),
            "start_time": "08:00", "end_time": "10:00",
            "session_type": "Cours",
        })
        other_date = (self.today + timedelta(days=1)).isoformat()
        resp = self.client.get(self._locations_url(), {
            "date": other_date,
            "start_time": "08:00", "end_time": "10:00",
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        room_ids = [r["id"] for r in resp.data["rooms"]]
        self.assertIn(self.room1.id, room_ids)
