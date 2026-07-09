import logging
from collections import defaultdict
from datetime import datetime
from io import BytesIO

import openpyxl
from django.db.models import Q
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill

from courses.models import Session
from users.models import Classe, Student
from users.permissions import apply_role_filtering

from .models import Attendance

logger = logging.getLogger(__name__)

STYLE_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
STYLE_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
STYLE_PRESENT_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
STYLE_ABSENT_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
STYLE_TITLE_FONT = Font(bold=True, size=12, color="1F4E79")
STYLE_BOLD = Font(bold=True)

DETAIL_HEADERS = [
    "Code Apogée",
    "Code Massar / CNE",
    "Nom",
    "Prénom",
    "Email",
    "Filière",
    "Classe",
    "Semestre",
    "Module",
    "Enseignant",
    "Date de la séance",
    "Heure de début",
    "Heure de fin",
    "Heure du scan",
    "Statut",
]

SUMMARY_HEADERS = [
    "Code Apogée",
    "Code Massar / CNE",
    "Nom",
    "Prénom",
    "Email",
    "Filière",
    "Classe",
    "Semestre",
    "Nombre total de séances",
    "Nombre de présences",
    "Nombre d'absences",
    "Taux de présence (%)",
]


class ExportWizardService:
    EXPORT_TYPE_ALL = "all"
    EXPORT_TYPE_PRESENT = "present"
    EXPORT_TYPE_ABSENT = "absent"

    EXPORT_MODE_BY_STUDENT = "by_student"
    EXPORT_MODE_BY_MODULE = "by_module"
    EXPORT_MODE_BY_SESSION = "by_session"
    EXPORT_MODE_RESUME_ABSENCES = "resume_absences"
    EXPORT_MODE_RESUME_PRESENCES = "resume_presences"

    @staticmethod
    def build_filtered_queryset(request, filters):
        qs = Attendance.objects.select_related(
            "student__classe__filiere",
            "student__classe__semestre",
            "session__teacher",
            "session__classe",
            "session__classe__filiere",
            "session__classe__semestre",
        )

        if request.user.role == "teacher":
            qs = qs.filter(session__teacher=request.user)
        elif request.user.role == "department_head":
            qs = apply_role_filtering(qs, "attendance", request)
        elif request.user.role == "program_manager":
            qs = apply_role_filtering(qs, "attendance", request)

        start_date = filters.get("start_date")
        if start_date:
            qs = qs.filter(session__date__gte=start_date)

        end_date = filters.get("end_date")
        if end_date:
            qs = qs.filter(session__date__lte=end_date)

        semestre = filters.get("semestre")
        if semestre:
            qs = qs.filter(student__classe__semestre__name__icontains=semestre)

        filiere = filters.get("filiere")
        if filiere:
            qs = qs.filter(student__classe__filiere__name__icontains=filiere)

        module = filters.get("module")
        if module:
            qs = qs.filter(session__subject__icontains=module)

        # ID-based hierarchical filters for professional export
        department_id = filters.get("department_id")
        if department_id:
            qs = qs.filter(student__classe__department_id=department_id)

        filiere_id = filters.get("filiere_id")
        if filiere_id:
            qs = qs.filter(student__classe__filiere_id=filiere_id)

        semestre_id = filters.get("semestre_id")
        if semestre_id:
            qs = qs.filter(student__classe__semestre_id=semestre_id)

        module_id = filters.get("module_id")
        if module_id:
            qs = qs.filter(session__module_id=module_id)

        classe_id = filters.get("classe_id")
        if classe_id:
            qs = qs.filter(session__classe_id=classe_id)

        export_type = filters.get("export_type", ExportWizardService.EXPORT_TYPE_ALL)
        if export_type == ExportWizardService.EXPORT_TYPE_PRESENT:
            qs = qs.filter(attendance_status=Attendance.STATUS_CONFIRMED)
        elif export_type == ExportWizardService.EXPORT_TYPE_ABSENT:
            qs = qs.filter(attendance_status=Attendance.STATUS_ABSENT)

        enseignant_id = filters.get("enseignant_id")
        if enseignant_id and request.user.role == "admin":
            qs = qs.filter(session__teacher_id=enseignant_id)

        etudiant_id = filters.get("etudiant_id")
        if etudiant_id:
            qs = qs.filter(student_id=etudiant_id)

        return qs

    @staticmethod
    def _build_detail_rows(qs):
        rows = []
        for att in qs.iterator(chunk_size=500):
            student = att.student
            session = att.session
            classe = student.classe

            local_time = timezone.localtime(att.validation_time) if att.validation_time else None
            statut = "Présent" if att.attendance_status == Attendance.STATUS_CONFIRMED else "Absent"

            rows.append({
                "code_apogee": student.code_apogee or "",
                "code_massar": student.code_massar or "",
                "nom": student.last_name,
                "prenom": student.first_name,
                "email": student.email or "",
                "filiere": classe.filiere.name if classe and classe.filiere else "",
                "classe": str(classe) if classe else "",
                "semestre": classe.semestre.name if classe and classe.semestre else "",
                "module": session.subject,
                "enseignant": session.teacher.full_name if session.teacher else "",
                "date_seance": str(session.date),
                "heure_debut": str(session.start_time) if session.start_time else "",
                "heure_fin": str(session.end_time) if session.end_time else "",
                "heure_scan": local_time.strftime("%H:%M:%S") if local_time else "",
                "statut": statut,
                "_student_id": student.id,
                "_session_id": session.id,
                "_module_name": session.subject,
            })
        return rows

    @staticmethod
    def _build_summary_rows(qs, export_type):
        students_data = defaultdict(lambda: {
            "total": 0, "presents": 0, "absents": 0,
        })

        for att in qs.iterator(chunk_size=500):
            student = att.student
            sid = student.id
            students_data[sid]["student"] = student
            students_data[sid]["total"] += 1

            if att.attendance_status == Attendance.STATUS_CONFIRMED:
                students_data[sid]["presents"] += 1
            elif att.attendance_status == Attendance.STATUS_ABSENT:
                students_data[sid]["absents"] += 1

        rows = []
        for sid, data in students_data.items():
            student = data["student"]
            classe = student.classe
            total = data["total"]
            presents = data["presents"]
            absents = data["absents"]
            taux = round((presents / total * 100), 1) if total else 0

            rows.append({
                "code_apogee": student.code_apogee or "",
                "code_massar": student.code_massar or "",
                "nom": student.last_name,
                "prenom": student.first_name,
                "email": student.email or "",
                "filiere": classe.filiere.name if classe and classe.filiere else "",
                "classe": str(classe) if classe else "",
                "semestre": classe.semestre.name if classe and classe.semestre else "",
                "total_seances": total,
                "presents": presents,
                "absents": absents,
                "taux_presence": f"{taux}%",
            })

        if export_type == ExportWizardService.EXPORT_MODE_RESUME_ABSENCES:
            rows.sort(key=lambda r: r["absents"], reverse=True)
        else:
            rows.sort(key=lambda r: r["taux_presence"], reverse=True)

        return rows

    @staticmethod
    def _style_header_row(ws, num_cols):
        for cell in ws[1]:
            cell.font = STYLE_HEADER_FONT
            cell.fill = STYLE_HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _auto_width(ws):
        for column in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in column), default=0)
            ws.column_dimensions[column[0].column_letter].width = min(max_len + 4, 40)

    @staticmethod
    def _compute_student_summary(detail_rows):
        total = len(detail_rows)
        presents = sum(1 for r in detail_rows if r["statut"] == "Présent")
        absents = total - presents
        taux = round((presents / total * 100), 1) if total else 0
        return {
            "total_seances": total,
            "presents": presents,
            "absents": absents,
            "taux_presence": f"{taux}%",
        }

    @staticmethod
    def _append_summary_to_sheet(ws, summary):
        ws.append([])
        ws.append(["RÉSUMÉ"])
        ws.cell(row=ws.max_row, column=1).font = STYLE_TITLE_FONT

        summary_headers = [
            "Nombre total de séances",
            "Nombre de présences",
            "Nombre d'absences",
            "Taux de présence (%)",
        ]
        summary_values = [
            summary["total_seances"],
            summary["presents"],
            summary["absents"],
            summary["taux_presence"],
        ]

        ws.append(summary_headers)
        for cell in ws[ws.max_row]:
            cell.font = STYLE_BOLD
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.append(summary_values)
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if summary["total_seances"] > 0:
                is_present = summary["presents"] >= summary["absents"]
                cell.fill = STYLE_PRESENT_FILL if is_present else STYLE_ABSENT_FILL

    @staticmethod
    def _generate_detail_xlsx(rows, export_mode, student_summary=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Présences"

        ws.append(DETAIL_HEADERS)
        ExportWizardService._style_header_row(ws, len(DETAIL_HEADERS))

        if export_mode == ExportWizardService.EXPORT_MODE_BY_STUDENT:
            rows.sort(key=lambda r: (r["nom"], r["prenom"], r["date_seance"]))
        elif export_mode == ExportWizardService.EXPORT_MODE_BY_MODULE:
            rows.sort(key=lambda r: (r["module"], r["date_seance"], r["nom"]))
        elif export_mode == ExportWizardService.EXPORT_MODE_BY_SESSION:
            rows.sort(key=lambda r: (r["date_seance"], r["module"], r["nom"]))

        for row in rows:
            ws.append([
                row["code_apogee"],
                row["code_massar"],
                row["nom"],
                row["prenom"],
                row["email"],
                row["filiere"],
                row["classe"],
                row["semestre"],
                row["module"],
                row["enseignant"],
                row["date_seance"],
                row["heure_debut"],
                row["heure_fin"],
                row["heure_scan"],
                row["statut"],
            ])
            last_row = ws.max_row
            is_present = row["statut"] == "Présent"
            for cell in ws[last_row]:
                cell.fill = STYLE_PRESENT_FILL if is_present else STYLE_ABSENT_FILL

        if student_summary:
            ExportWizardService._append_summary_to_sheet(ws, student_summary)

        ExportWizardService._auto_width(ws)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()

    @staticmethod
    def _generate_summary_xlsx(rows, export_mode):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Résumé des présences"

        ws.append(SUMMARY_HEADERS)
        ExportWizardService._style_header_row(ws, len(SUMMARY_HEADERS))

        for row in rows:
            ws.append([
                row["code_apogee"],
                row["code_massar"],
                row["nom"],
                row["prenom"],
                row["email"],
                row["filiere"],
                row["classe"],
                row["semestre"],
                row["total_seances"],
                row["presents"],
                row["absents"],
                row["taux_presence"],
            ])

        ExportWizardService._auto_width(ws)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()

    @staticmethod
    def _build_professional_rows(qs):
        """Build per-student summary with department, module, teacher info."""
        from collections import defaultdict
        students_data = defaultdict(lambda: {
            "total": 0, "presents": 0, "absents": 0,
        })

        for att in qs.iterator(chunk_size=500):
            student = att.student
            sid = student.id
            d = students_data[sid]
            d["student"] = student
            d["total"] += 1
            if att.attendance_status == Attendance.STATUS_CONFIRMED:
                d["presents"] += 1
            elif att.attendance_status == Attendance.STATUS_ABSENT:
                d["absents"] += 1
            d["session"] = att.session

        rows = []
        for sid, data in students_data.items():
            student = data["student"]
            classe = student.classe
            session = data.get("session")
            total = data["total"]
            presents = data["presents"]
            absents = data["absents"]
            taux_presence = round((presents / total * 100), 1) if total else 0
            taux_absence = round((absents / total * 100), 1) if total else 0

            # Get department: prefer classe.department, fallback to filiere.department
            department_name = ""
            if classe and classe.department:
                department_name = classe.department.name
            elif classe and classe.filiere and classe.filiere.department:
                department_name = classe.filiere.department.name

            module_name = session.module.name if session and session.module else (session.subject if session else "")
            teacher_name = session.teacher.full_name if session and session.teacher else ""

            rows.append({
                "code_massar": student.code_massar or "",
                "code_apogee": student.code_apogee or "",
                "nom": student.last_name,
                "prenom": student.first_name,
                "department": department_name,
                "filiere": classe.filiere.name if classe and classe.filiere else "",
                "semestre": classe.semestre.name if classe and classe.semestre else "",
                "module": module_name,
                "classe": str(classe) if classe else "",
                "enseignant": teacher_name,
                "total_seances": total,
                "presents": presents,
                "absents": absents,
                "taux_presence": f"{taux_presence}%",
                "taux_absence": f"{taux_absence}%",
            })

        rows.sort(key=lambda r: r["absents"], reverse=True)
        return rows

    @staticmethod
    def _generate_professional_xlsx(rows, filters=None):
        """Generate a professional XLSX with header, stats, and formatted table."""
        from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rapport d'absences"

        # Colors
        DARK_BLUE = "1F4E79"
        LIGHT_BLUE = "D6E4F0"
        WHITE = "FFFFFF"
        LIGHT_GRAY = "F2F2F2"
        GREEN = "C6EFCE"
        RED = "FFC7CE"

        header_fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
        header_font = Font(bold=True, color=WHITE, size=10)
        title_font = Font(bold=True, size=14, color=DARK_BLUE)
        subtitle_font = Font(bold=True, size=11, color=DARK_BLUE)
        stat_label_font = Font(bold=True, size=10, color=DARK_BLUE)
        stat_value_font = Font(size=10)
        thin_border = Border(
            left=Side(style="thin", color="999999"),
            right=Side(style="thin", color="999999"),
            top=Side(style="thin", color="999999"),
            bottom=Side(style="thin", color="999999"),
        )

        # === HEADER SECTION ===
        # Row 1: Title (merged across all columns)
        LAST_COL_LETTER = "O"
        ws.merge_cells(f"A1:{LAST_COL_LETTER}1")
        title_cell = ws["A1"]
        title_cell.value = "Rapport des absences"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Row 2: University / Faculty / Dept / Filiere / Semestre / Module / Classe / Enseignant
        ws.merge_cells(f"A2:{LAST_COL_LETTER}2")
        sub_cell = ws["A2"]
        dept_name = ""
        filiere_name = ""
        semestre_name = ""
        module_name = ""
        classe_name = ""
        enseignant_name = ""
        if filters:
            dept_name = filters.get("department_name", "")
            filiere_name = filters.get("filiere_name", "")
            semestre_name = filters.get("semestre_name", "")
            module_name = filters.get("module_name", "")
            classe_name = filters.get("classe_name", "")
            enseignant_name = filters.get("enseignant_name", "")
        header_parts = ["Université", "Faculté des Sciences"]
        if dept_name:
            header_parts.append(f"Département: {dept_name}")
        if filiere_name:
            header_parts.append(f"Filière: {filiere_name}")
        if semestre_name:
            header_parts.append(f"Semestre: {semestre_name}")
        if module_name:
            header_parts.append(f"Module: {module_name}")
        if classe_name:
            header_parts.append(f"Classe: {classe_name}")
        if enseignant_name:
            header_parts.append(f"Enseignant: {enseignant_name}")
        sub_cell.value = " | ".join(header_parts)
        sub_cell.font = Font(size=10, italic=True, color="555555")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        # Row 3: Period / Academic year / Generation date / Generated by
        ws.merge_cells(f"A3:{LAST_COL_LETTER}3")
        period_cell = ws["A3"]
        start_date = filters.get("start_date", "") if filters else ""
        end_date = filters.get("end_date", "") if filters else ""
        period_text = f"{start_date} → {end_date}" if start_date and end_date else "Toute la période"
        academic_year = filters.get("academic_year", "") if filters else ""
        gen_date = datetime.now().strftime("%d/%m/%Y %H:%M")
        gen_by = filters.get("generated_by", "") if filters else ""
        period_parts = [f"Période: {period_text}"]
        if academic_year:
            period_parts.append(f"Année univ.: {academic_year}")
        period_parts.append(f"Généré le: {gen_date}")
        if gen_by:
            period_parts.append(f"Par: {gen_by}")
        period_cell.value = " | ".join(period_parts)
        period_cell.font = Font(size=9, italic=True, color="777777")
        period_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 20

        # Row 4: Empty spacer
        ws.row_dimensions[4].height = 6

        # === STATS SECTION (rows 5-6) ===
        stats_start = 5
        total_students = len(rows)
        total_presents = sum(r["presents"] for r in rows)
        total_absents = sum(r["absents"] for r in rows)
        total_presence_rate = round((total_presents / (total_presents + total_absents) * 100), 1) if (total_presents + total_absents) else 0
        total_absence_rate = round((total_absents / (total_presents + total_absents) * 100), 1) if (total_presents + total_absents) else 0

        stats_data = [
            ("Total étudiants", str(total_students)),
            ("Total présences", str(total_presents)),
            ("Total absences", str(total_absents)),
            ("Taux de présence", f"{total_presence_rate}%"),
            ("Taux d'absence", f"{total_absence_rate}%"),
        ]

        # Write stats in two rows: labels in row 5, values in row 6
        col = 1
        for label, value in stats_data:
            cell_label = ws.cell(row=stats_start, column=col, value=label)
            cell_label.font = stat_label_font
            cell_label.alignment = Alignment(horizontal="center", vertical="center")
            cell_label.fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")
            cell_label.border = thin_border

            cell_value = ws.cell(row=stats_start + 1, column=col, value=value)
            cell_value.font = stat_value_font
            cell_value.alignment = Alignment(horizontal="center", vertical="center")
            cell_value.border = thin_border
            col += 1

        # Row 7: Empty spacer
        ws.row_dimensions[stats_start + 2].height = 6

        # === DATA TABLE ===
        data_start = stats_start + 3
        PROFESSIONAL_HEADERS = [
            "Code Massar", "Code Apogée", "Nom", "Prénom",
            "Département", "Filière", "Semestre", "Module",
            "Classe", "Enseignant",
            "Total séances", "Présences", "Absences", "Taux présence", "Taux absence",
        ]
        NUM_COLS = len(PROFESSIONAL_HEADERS)
        HEADER_ROW = data_start

        for ci, header in enumerate(PROFESSIONAL_HEADERS, 1):
            cell = ws.cell(row=HEADER_ROW, column=ci, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[HEADER_ROW].height = 22

        # Data rows
        for ri, row in enumerate(rows):
            excel_row = HEADER_ROW + 1 + ri
            values = [
                row.get("code_massar", ""),
                row.get("code_apogee", ""),
                row.get("nom", ""),
                row.get("prenom", ""),
                row.get("department", ""),
                row.get("filiere", ""),
                row.get("semestre", ""),
                row.get("module", ""),
                row.get("classe", ""),
                row.get("enseignant", ""),
                row.get("total_seances", 0),
                row.get("presents", 0),
                row.get("absents", 0),
                row.get("taux_presence", ""),
                row.get("taux_absence", ""),
            ]
            is_absent = row.get("absents", 0) > row.get("presents", 0)

            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=excel_row, column=ci, value=val)
                cell.font = Font(size=9)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
                # Alternating row colors
                if ri % 2 == 1:
                    cell.fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

            # Highlight absent students row
            if is_absent:
                for ci in range(1, NUM_COLS + 1):
                    cell = ws.cell(row=excel_row, column=ci)
                    cell.fill = PatternFill(start_color=RED, end_color=RED, fill_type="solid")

        # Auto-width
        for ci in range(1, NUM_COLS + 1):
            max_len = len(PROFESSIONAL_HEADERS[ci - 1])
            for ri in range(HEADER_ROW + 1, HEADER_ROW + 1 + len(rows)):
                cell_val = ws.cell(row=ri, column=ci).value
                if cell_val:
                    max_len = max(max_len, len(str(cell_val)))
            ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 3, 35)

        # Freeze pane at header row
        ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

        # Auto-filter
        last_data_row = HEADER_ROW + len(rows)
        if rows:
            ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(NUM_COLS)}{last_data_row}"

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()

    @staticmethod
    def generate_xlsx(request, filters):
        export_mode = filters.get("export_mode", ExportWizardService.EXPORT_MODE_BY_STUDENT)
        etudiant_id = filters.get("etudiant_id")

        qs = ExportWizardService.build_filtered_queryset(request, filters)

        if export_mode == "professional":
            if filters is None:
                filters = {}
            filters["generated_by"] = request.user.get_full_name() or request.user.email
            rows = ExportWizardService._build_professional_rows(qs)
            content = ExportWizardService._generate_professional_xlsx(rows, filters)
            filename = "rapport_absences.xlsx"
        elif export_mode in (
            ExportWizardService.EXPORT_MODE_RESUME_ABSENCES,
            ExportWizardService.EXPORT_MODE_RESUME_PRESENCES,
        ):
            rows = ExportWizardService._build_summary_rows(qs, export_mode)
            content = ExportWizardService._generate_summary_xlsx(rows, export_mode)
            filename = "resume_presences.xlsx"
        else:
            rows = ExportWizardService._build_detail_rows(qs)
            student_summary = None
            if etudiant_id:
                student_summary = ExportWizardService._compute_student_summary(rows)
            content = ExportWizardService._generate_detail_xlsx(rows, export_mode, student_summary)
            filename = "presences.xlsx"

        return content, filename
